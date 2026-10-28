"""Elo por superficie: lo calcula el Estadístico y lo usan la sala y el Especialista en superficies.

Se parte de los resultados públicos de tennis-data.co.uk (ATP y WTA de las tres últimas temporadas, con la
superficie de cada partido) y se añaden los resultados que la oficina va recogiendo cada día. Si tennis-data no
responde, se prueba con los ficheros de Jeff Sackmann en GitHub. Se recalcula entero una vez al día.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, timedelta

import requests

from agentes import comun, config, cuotas

INICIAL = 1500.0
SUPERFICIES = ("hard", "clay", "grass")
TENNIS_DATA = "http://www.tennis-data.co.uk/{anio}{w}/{anio}.xlsx"
SACKMANN = "https://raw.githubusercontent.com/JeffSackmann/tennis_{c}/master/{c}_matches_{anio}.csv"


def k_factor(partidos: int) -> float:
    """Como el Elo de FiveThirtyEight: los jugadores con pocos partidos se mueven más."""
    return 250 / (partidos + 5) ** 0.4


def esperado(a: float, b: float) -> float:
    return 1 / (1 + 10 ** ((b - a) / 400))


def _superficie(texto: str) -> str:
    t = (texto or "").lower()
    return "clay" if "clay" in t else "grass" if "grass" in t else "hard"


def _fecha(valor) -> str:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    t = str(valor or "")
    if len(t) == 8 and t.isdigit():
        return f"{t[:4]}-{t[4:6]}-{t[6:]}"
    return t[:10]


def leer_tennis_data(contenido: bytes, circuito: str) -> list[dict]:
    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
    filas = hoja.iter_rows(values_only=True)
    cab = [str(c or "").strip() for c in next(filas)]
    i = {c: cab.index(c) for c in ("Date", "Surface", "Winner", "Loser", "Tournament") if c in cab}
    com = cab.index("Comment") if "Comment" in cab else None
    partidos = []
    for f in filas:
        if not f or not f[i["Winner"]] or not f[i["Loser"]]:
            continue
        if com is not None and str(f[com] or "").lower().startswith("walk"):
            continue
        partidos.append({"fecha": _fecha(f[i["Date"]]), "superficie": _superficie(f[i["Surface"]]),
                         "ganador": comun.clave_abreviada(str(f[i["Winner"]])), "perdedor": comun.clave_abreviada(str(f[i["Loser"]])),
                         "nombres": [str(f[i["Winner"]]), str(f[i["Loser"]])], "circuito": circuito})
    return partidos


def leer_sackmann(texto: str, circuito: str) -> list[dict]:
    partidos = []
    for f in csv.DictReader(io.StringIO(texto)):
        if not f.get("winner_name") or not f.get("loser_name"):
            continue
        partidos.append({"fecha": _fecha(f.get("tourney_date")), "superficie": _superficie(f.get("surface")),
                         "ganador": comun.claves_completas(f["winner_name"])[0], "perdedor": comun.claves_completas(f["loser_name"])[0],
                         "nombres": [f["winner_name"], f["loser_name"]], "circuito": circuito})
    return partidos


def descargar_base(anio: int, sesion=None) -> list[dict]:
    sesion = sesion or requests
    partidos = []
    for circuito, w in (("ATP", ""), ("WTA", "w")):
        leidos = []
        for a in (anio - 2, anio - 1, anio):
            try:
                r = sesion.get(TENNIS_DATA.format(anio=a, w=w), timeout=60)
                r.raise_for_status()
                leidos += leer_tennis_data(r.content, circuito)
            except Exception as e:  # noqa: BLE001 - cualquier fallo de una fuente se salta
                print(f"tennis-data {circuito} {a}: {e}")
        if not leidos:
            for a in (anio - 2, anio - 1, anio):
                try:
                    r = sesion.get(SACKMANN.format(c=circuito.lower(), anio=a), timeout=60)
                    r.raise_for_status()
                    leidos += leer_sackmann(r.text, circuito)
                except Exception as e:  # noqa: BLE001
                    print(f"Sackmann {circuito} {a}: {e}")
        print(f"Elo: {len(leidos)} partidos de la {circuito}")
        partidos += leidos
    return partidos


def de_resultados(resultados: list[dict], tabla: dict) -> list[dict]:
    """Los resultados propios (nombre completo) con la clave con que ya aparece cada jugador, si aparece."""
    partidos = []
    for r in resultados:
        if not r.get("completo") or r.get("ganador") is None:
            continue
        nombres = [r["jugadores"][r["ganador"]], r["jugadores"][1 - r["ganador"]]]
        claves = [comun.buscar_jugador(n, tabla) or comun.claves_completas(n)[0] for n in nombres]
        partidos.append({"fecha": r.get("fecha", ""), "superficie": cuotas.superficie(r.get("torneo", "")),
                         "ganador": claves[0], "perdedor": claves[1], "nombres": nombres, "circuito": r.get("circuito", "")})
    return partidos


def calcular(partidos: list[dict]) -> dict:
    jugadores: dict[str, dict] = {}
    h2h: dict[str, list[int]] = {}
    vistos = set()
    for p in sorted(partidos, key=lambda x: x["fecha"]):
        par = "~".join(sorted((p["ganador"], p["perdedor"])))
        try:
            dia = date.fromisoformat(p["fecha"][:10])
            cerca = {(str(dia + timedelta(days=d)), par) for d in (-1, 0, 1)}
        except ValueError:
            cerca = {(p["fecha"], par)}
        if cerca & vistos:  # el mismo partido de dos fuentes (las fechas pueden ir un día movidas)
            continue
        vistos.add((p["fecha"][:10], par))
        g, pe = (jugadores.setdefault(k, {"nombre": n, "circuito": p.get("circuito", ""),
                                          "elo": {"all": INICIAL, **{s: INICIAL for s in SUPERFICIES}},
                                          "n": {"all": 0, **{s: 0 for s in SUPERFICIES}}, "ultimos": []})
                 for k, n in ((p["ganador"], p["nombres"][0]), (p["perdedor"], p["nombres"][1])))
        s = p["superficie"]
        for tipo in ("all", s):
            e = esperado(g["elo"][tipo], pe["elo"][tipo])
            kg, kp = k_factor(g["n"][tipo]), k_factor(pe["n"][tipo])
            g["elo"][tipo] += kg * (1 - e)
            pe["elo"][tipo] -= kp * (1 - e)
            g["n"][tipo] += 1
            pe["n"][tipo] += 1
        for j, r, rival in ((g, "G", pe), (pe, "P", g)):
            j["ultimos"] = (j["ultimos"] + [f"{p['fecha']} {r} {s} vs {rival['nombre']}"])[-10:]
        h2h.setdefault(par, [0, 0])[0 if p["ganador"] == par.split("~")[0] else 1] += 1
    for j in jugadores.values():
        j["elo"] = {k: round(v, 1) for k, v in j["elo"].items()}
    return {"actualizado": comun.iso(comun.ahora()), "jugadores": jugadores, "h2h": h2h}


def actualizar(sesion=None, forzar: bool = False) -> dict:
    """Recalcula el Elo (una vez al día salvo que se fuerce)."""
    previo = comun.leer(config.ELO, {})
    if not forzar and previo.get("actualizado", "")[:10] == comun.ahora().date().isoformat():
        return previo
    base = descargar_base(comun.ahora().year, sesion)
    if not base and previo.get("jugadores"):
        print("Elo: sin fuentes hoy, se mantiene el de ayer")
        return previo
    tabla = {p["ganador"]: 1 for p in base} | {p["perdedor"]: 1 for p in base}
    elo = calcular(base + de_resultados(comun.leer(config.RESULTADOS, []), tabla))
    comun.guardar(config.ELO, elo)
    return elo


def probabilidad(elo: dict, a: str, b: str, superficie: str) -> dict | None:
    """Probabilidad de que gane `a`: mezcla del Elo general y el de la superficie (este pesa más cuantos más
    partidos haya en ella)."""
    t = elo.get("jugadores", {})
    ka, kb = comun.buscar_jugador(a, t), comun.buscar_jugador(b, t)
    if not ka or not kb:
        return None
    ja, jb = t[ka], t[kb]
    general = esperado(ja["elo"]["all"], jb["elo"]["all"])
    s = superficie if superficie in SUPERFICIES else "hard"
    en_sup = esperado(ja["elo"][s], jb["elo"][s])
    peso = 0.5 * min(1.0, min(ja["n"][s], jb["n"][s]) / 20)
    par = "~".join(sorted((ka, kb)))
    cara = elo.get("h2h", {}).get(par, [0, 0])
    if par.split("~")[0] != ka:
        cara = cara[::-1]
    return {"prob": (1 - peso) * general + peso * en_sup, "general": general, "superficie": en_sup,
            "elo_a": ja["elo"], "elo_b": jb["elo"], "n_a": ja["n"], "n_b": jb["n"],
            "forma_a": ja["ultimos"][-5:], "forma_b": jb["ultimos"][-5:], "h2h": cara}
