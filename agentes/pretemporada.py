"""Pretemporada simulada: los analistas repasan día a día unos meses ya jugados para llegar entrenados.

Cada día de la simulación se hace lo mismo que en la oficina, pero con datos del pasado y sin gastar créditos de
cuotas: la cartelera sale de tennis-data.co.uk (partidos reales con la cuota de Bet365, o la media del mercado, y
la de Pinnacle), la sala escribe su informe, los analistas deciden (con la apuesta de práctica), por la noche se
liquida con el resultado real y cada uno escribe su lección; los domingos, la reflexión y la biblioteca.
No hay noticias de esos días ni despidos. El Elo y el cansancio de cada día solo usan partidos anteriores.

Se guarda en datos/pretemporada (en GitHub, en una rama aparte para no publicar la web cada vez) y avanza unos días
en cada ejecución. Al terminar, `--trasladar` pasa a la oficina las reglas, lecciones, reflexiones y la
biblioteca. El dinero de la oficina no se toca.

    python -m agentes.pretemporada --dias 8          # avanza 8 días
    python -m agentes.pretemporada --trasladar       # al terminar: lo aprendido pasa a la oficina
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import io
from datetime import date, datetime, time, timedelta

import requests

from agentes import aprendizaje, boletos, ciclo, comun, config, elo as elo_mod, emociones, ia as ia_mod, oficina, prompts, resultados, sala

DESDE, HASTA = "2026-07-01", "2026-10-28"
MAX_PARTIDOS = 16            # como mucho 16 partidos al día (los de mejor ranking), como una cartelera normal
CASAS = (("B365", "bet365", "Bet365"), ("Avg", "media", "Media del mercado"))
NOTA = ("PRETEMPORADA: estás entrenando con días ya jugados. Decide igual que si fuera hoy. De esos días no hay "
        "noticias, solo hay cuotas del ganador y la cuota de Pinnacle es la de cierre.")


# --- Datos de tennis-data --------------------------------------------------------------------------
def nombre_completo(nombre: str) -> str:
    """«De Minaur A.» → «A. De Minaur» (así lo reconocen el Elo y los resultados)."""
    clave = comun.clave_abreviada(nombre)
    partes = nombre.strip().split()
    if "|" not in clave or len(partes) < 2 or not partes[-1].endswith("."):
        return nombre.strip()
    iniciales = []
    while len(partes) > 1 and partes[-1].endswith("."):
        iniciales.insert(0, partes.pop())
    return " ".join(iniciales + partes)


def _num(valor) -> float | None:
    try:
        x = float(valor)
    except (TypeError, ValueError):
        return None
    return x if x > 1 else None


def leer_dias(contenido: bytes, circuito: str) -> list[dict]:
    """Los partidos de un fichero de tennis-data con su resultado (sets) y sus cuotas."""
    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
    filas = hoja.iter_rows(values_only=True)
    cab = [str(c or "").strip() for c in next(filas)]
    col = {c: i for i, c in enumerate(cab)}
    v = lambda f, c: f[col[c]] if c in col and col[c] < len(f) else None  # noqa: E731
    partidos = []
    for f in filas:
        if not f or not v(f, "Winner") or not v(f, "Loser"):
            continue
        w, l = nombre_completo(str(v(f, "Winner"))), nombre_completo(str(v(f, "Loser")))
        fecha = elo_mod._fecha(v(f, "Date"))
        sets = []
        for n in range(1, 6):
            a, b = _num_juegos(v(f, f"W{n}")), _num_juegos(v(f, f"L{n}"))
            if a is not None and b is not None:
                sets.append([a, b])
        comentario = str(v(f, "Comment") or "Completed")
        casa = next(((k, nombre) for pre, k, nombre in CASAS if _num(v(f, f"{pre}W")) and _num(v(f, f"{pre}L"))), None)
        rangos = [_num_juegos(v(f, "WRank")), _num_juegos(v(f, "LRank"))]
        partidos.append({
            "fecha": fecha, "torneo": str(v(f, "Tournament") or ""), "circuito": circuito,
            "superficie": elo_mod._superficie(str(v(f, "Surface") or "")), "ganador": w, "perdedor": l,
            "sets": sets, "completo": comentario.lower().startswith("complet"), "comentario": comentario,
            "casa": casa[0] if casa else None,
            "cuotas": [_num(v(f, f"{_pre(casa[0])}W")), _num(v(f, f"{_pre(casa[0])}L"))] if casa else None,
            "pinnacle": [_num(v(f, "PSW")), _num(v(f, "PSL"))],
            "ranking": sum(r if r else 2000 for r in rangos)})
    return partidos


def _pre(casa: str) -> str:
    return next(pre for pre, k, _ in CASAS if k == casa)


def _num_juegos(valor) -> int | None:
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return None


def descargar(anios: tuple[int, ...], sesion=None) -> tuple[list[dict], list[dict]]:
    """(partidos para el Elo de los tres últimos años, partidos con cuotas y resultado de `anios`)."""
    sesion = sesion or requests
    base, dias = [], []
    for circuito, w in (("ATP", ""), ("WTA", "w")):
        for a in range(min(anios) - 2, max(anios) + 1):
            try:
                r = sesion.get(elo_mod.TENNIS_DATA.format(anio=a, w=w), headers=elo_mod.CABECERAS, timeout=60)
                r.raise_for_status()
            except Exception as e:  # noqa: BLE001
                print(f"tennis-data {circuito} {a}: {e}")
                continue
            base += elo_mod.leer_tennis_data(r.content, circuito)
            if a in anios:
                dias += leer_dias(r.content, circuito)
    print(f"Pretemporada: {len(base)} partidos para el Elo y {len(dias)} con resultado")
    return base, dias


# --- Un día de la simulación -----------------------------------------------------------------------
def _id(p: dict) -> str:
    return hashlib.sha1(f"{p['fecha']}|{p['ganador']}|{p['perdedor']}".encode()).hexdigest()[:32]


def cartelera(del_dia: list[dict], dia: date) -> tuple[list[dict], list[dict]]:
    """(cartelera del día como la de la oficina, resultados de esos partidos). El ganador no va siempre primero."""
    elegidos = sorted([p for p in del_dia if p["cuotas"]], key=lambda p: (p["ranking"], p["ganador"]))[:MAX_PARTIDOS]
    partidos, res = [], []
    for i, p in enumerate(sorted(elegidos, key=lambda p: (p["torneo"], p["ganador"]))):
        pid = _id(p)
        al_reves = int(pid, 16) % 2 == 1
        jug = [p["perdedor"], p["ganador"]] if al_reves else [p["ganador"], p["perdedor"]]
        cuotas = p["cuotas"][::-1] if al_reves else p["cuotas"]
        pin = p["pinnacle"][::-1] if al_reves else p["pinnacle"]
        inicio = datetime.combine(dia, time(11, 0), tzinfo=comun.MADRID) + timedelta(minutes=20 * i)
        opciones = [{"id": f"{pid}|h2h|{j}|", "mercado": "h2h", "nombre": j, "punto": None, "cuota": c, "pinnacle": q}
                    for j, c, q in zip(jug, cuotas, pin)]
        partidos.append({"id": pid, "deporte": f"sim_{p['circuito'].lower()}", "torneo": p["torneo"],
                         "circuito": p["circuito"], "superficie": p["superficie"], "inicio": comun.iso(inicio),
                         "jugadores": jug, "opciones": opciones, "casa": p["casa"],
                         "pinnacle_h2h": {j: q for j, q in zip(jug, pin) if q}})
        sets = [s[::-1] for s in p["sets"]] if al_reves else p["sets"]
        res.append({"id": pid, "fecha": p["fecha"], "torneo": p["torneo"], "circuito": p["circuito"], "jugadores": jug,
                    "ganador": 1 if al_reves else 0, "sets": sets, "completo": p["completo"], "estado": p["comentario"]})
    return partidos, res


class Contador(ia_mod.IA):
    """Cuenta las llamadas y los fallos: si la IA falla (sin cuota), el día se repite en la siguiente ejecución."""

    def __init__(self, ia):
        self.ia, self.llamadas, self.fallos = ia, 0, 0

    def pedir_json(self, instrucciones, mensaje, esquema=None):
        self.llamadas += 1
        try:
            return self.ia.pedir_json(instrucciones, mensaje, esquema)
        except Exception:
            self.fallos += 1
            raise


def simular_dia(estado: dict, apuestas: list[dict], dia: date, base: list[dict], todos: list[dict], ia) -> int:
    """Un día entero: cartelera, sala, decisiones, liquidación, lecciones y (en domingo) reflexiones."""
    hoy = dia.isoformat()
    previos = [p for p in base if p["fecha"] < hoy]
    historial = [r for p in todos if p["fecha"] < hoy
                 for r in [{"fecha": p["fecha"], "torneo": p["torneo"], "jugadores": [p["ganador"], p["perdedor"]],
                            "ganador": 0, "sets": p["sets"], "completo": p["completo"]}]]
    partidos, res = cartelera([p for p in todos if p["fecha"] == hoy], dia)
    if not partidos:
        return 0
    original = resultados.historial
    resultados.historial = lambda: {"dias": [], "partidos": historial}  # el cansancio solo ve el pasado
    try:
        _jornada(estado, apuestas, dia, partidos, res, elo_mod.calcular(previos), ia)
    finally:
        resultados.historial = original
    return len(partidos)


def _jornada(estado: dict, apuestas: list[dict], dia: date, partidos: list[dict], res: list[dict], elo: dict, ia) -> None:
    hoy = dia.isoformat()
    manana = datetime.combine(dia, time(9, 0), tzinfo=comun.MADRID)
    noche = datetime.combine(dia, time(23, 30), tzinfo=comun.MADRID)
    estado["ultima_decision"] = hoy
    estado.setdefault("emociones_desde", comun.iso(manana))
    aprendizaje.apuntar_tonta(estado, partidos, manana)
    sala.escribir(estado, partidos, elo, {"noticias": []}, ia, manana)
    sala_txt = prompts.texto_sala(sala.para_decidir(estado, manana))
    ciclo.decidir_todos(estado, apuestas, partidos, elo, {"noticias": []}, sala_txt, ia, manana, NOTA, practica=True)
    # Por la noche: resultados reales
    agentes = {a["id"]: a for a in estado["agentes"]}
    por_id = {r["id"]: r for r in res}
    pendientes = [a for a in apuestas if a["estado"] == "pendiente"]
    for a in pendientes:
        r = por_id.get(a["partido"])
        a["pinnacle_cierre"] = a.get("pinnacle_apuesta")  # tennis-data da la de cierre
        puntos = boletos.puntuar(a, r) if r else 0.0
        boletos.liquidar(a, agentes[a["agente"]], 0.0 if puntos is None else puntos,
                         resultados.texto(r) if r else "Sin resultado: anulada", noche)
    for agente in oficina.activos(estado):
        emociones.por_liquidadas(agente, [a for a in pendientes if a["agente"] == agente["id"]], noche)
        emociones.actualizar(agente, estado, apuestas, noche)
    aprendizaje.puntuar_sala(estado, res, noche)
    aprendizaje.puntuar_descartes(estado, res, noche)
    aprendizaje.liquidar_tonta(estado, res, noche)
    oficina.lecciones(estado, apuestas, ia, noche)
    if dia.weekday() == 6:
        oficina.reflexiones(estado, apuestas, ia, noche)
        oficina.ordenar_biblioteca(estado, apuestas, ia, noche)
    nuevas = [a for a in apuestas if a["fecha"] == hoy]
    oficina.anotar(estado, noche, f"🏋️ {hoy}: {len(partidos)} partidos, {len(nuevas)} apuestas "
                                  f"({sum(a['ganancia'] or 0 for a in nuevas):+.2f} €).")


# --- Avanzar y trasladar ---------------------------------------------------------------------------
def rutas() -> dict:
    d = config.DATOS / "pretemporada"
    return {"estado": d / "estado.json", "apuestas": d / "apuestas.json", "progreso": d / "progreso.json"}


def avanzar(dias: int, desde: str = DESDE, hasta: str = HASTA, ia=None, sesion=None, activar: bool = False,
            solo_si_activa: bool = False) -> dict:
    f = rutas()
    progreso = comun.leer(f["progreso"], None) or {"desde": desde, "hasta": hasta, "siguiente": desde, "dias": 0,
                                                    "llamadas": 0, "terminada": False, "trasladada": False, "activa": False}
    if activar:
        progreso["activa"] = True
    if progreso["terminada"]:
        print("La pretemporada ya está terminada.")
        return progreso
    if solo_si_activa and not progreso.get("activa"):
        print("La pretemporada no está activada: solo avanza a mano (o con --activar).")
        comun.guardar(f["progreso"], progreso)
        return progreso
    inicio = datetime.combine(date.fromisoformat(progreso["desde"]), time(9, 0), tzinfo=comun.MADRID)
    estado = comun.leer(f["estado"], None) or oficina.estado_inicial(inicio, "pretemporada")
    estado["casa"] = "bet365"
    apuestas = comun.leer(f["apuestas"], [])
    base, todos = descargar(tuple(range(int(progreso["siguiente"][:4]), int(progreso["hasta"][:4]) + 1)), sesion)
    if not todos:
        print("Sin datos de tennis-data: se prueba en la siguiente ejecución.")
        return progreso
    if ia is None:
        ia = ia_mod.Claude() if config.PROVEEDOR_IA == "claude" else ia_mod.Gemini(config.MODELO_PRETEMPORADA, respaldo=None)
    contador = Contador(ia)
    con_partidos = sorted({p["fecha"] for p in todos if p["cuotas"]})
    hechos = 0
    dia = date.fromisoformat(progreso["siguiente"])
    fin = date.fromisoformat(progreso["hasta"])
    while dia <= fin and hechos < dias:
        if dia.isoformat() not in con_partidos:
            if dia.isoformat() > con_partidos[-1] and comun.ahora().date() - dia < timedelta(days=10):
                print(f"tennis-data aún no tiene partidos del {dia}: se sigue más adelante.")  # se actualiza cada semana
                break
            dia += timedelta(days=1)
            continue
        copia, copia_apuestas, antes = copy.deepcopy(estado), copy.deepcopy(apuestas), contador.fallos
        simular_dia(estado, apuestas, dia, base, todos, contador)
        if contador.fallos > antes:  # la IA no respondió (sin cuota): se repite el día otro rato
            estado, apuestas = copia, copia_apuestas
            print(f"La IA falló el {dia}: se para aquí y se repite en la siguiente ejecución.")
            break
        hechos += 1
        progreso["dias"] += 1
        dia += timedelta(days=1)
    else:
        if dia > fin:
            progreso["terminada"] = True
    progreso["siguiente"] = dia.isoformat()
    progreso["llamadas"] += contador.llamadas
    progreso["resumen"] = resumen(estado, apuestas)
    comun.guardar(f["estado"], estado)
    comun.guardar(f["apuestas"], apuestas)
    comun.guardar(f["progreso"], progreso)
    print(f"Pretemporada: {hechos} días hechos (van {progreso['dias']}), siguiente {progreso['siguiente']}, "
          f"{contador.llamadas} llamadas a la IA" + (" · TERMINADA" if progreso["terminada"] else ""))
    return progreso


def resumen(estado: dict, apuestas: list[dict]) -> dict:
    return {"agentes": {a["id"]: {"nombre": a["nombre"], **boletos.resumen([x for x in apuestas if x["agente"] == a["id"]]),
                                  "practica": sum(1 for x in apuestas if x["agente"] == a["id"] and x.get("practica"))}
                        for a in estado["agentes"]},
            "tonta": aprendizaje.resumen_tonta(estado), "sala": estado.get("sala_aciertos")}


def trasladar(momento: datetime | None = None) -> bool:
    """Pasa a la oficina lo aprendido: reglas, lecciones, reflexiones, biblioteca y el resumen de cada uno."""
    f = rutas()
    progreso = comun.leer(f["progreso"], None)
    if not progreso or not progreso.get("terminada") or progreso.get("trasladada"):
        print("No hay una pretemporada terminada por trasladar.")
        return False
    momento = momento or comun.ahora()
    sim, apuestas_sim = comun.leer(f["estado"], {}), comun.leer(f["apuestas"], [])
    estado = oficina.cargar(momento)
    por_rol = {a["rol"]: a for a in sim.get("agentes", [])}
    for agente in oficina.activos(estado):
        s = por_rol.get(agente["rol"])
        if not s:
            continue
        if not agente.get("reglas"):
            agente["reglas"] = s.get("reglas", [])
            agente["reglas_retiradas"] = s.get("reglas_retiradas", [])
        agente["lecciones"] = (s.get("lecciones", []) + agente.get("lecciones", []))[-aprendizaje.MAX_LECCIONES:]
        agente["reflexiones"] = (s.get("reflexiones", []) + agente.get("reflexiones", []))[-oficina.MAX_REFLEXIONES:]
        suyas = [a for a in apuestas_sim if a["agente"] == s["id"]]
        agente["pretemporada"] = {"desde": progreso["desde"], "hasta": progreso["hasta"], "resumen": boletos.resumen(suyas),
                                  "desglose": aprendizaje.desglose(suyas)}
    if not estado.get("biblioteca"):
        estado["biblioteca"] = sim.get("biblioteca", [])
    estado["pretemporada"] = {"desde": progreso["desde"], "hasta": progreso["hasta"], "dias": progreso["dias"],
                              **resumen(sim, apuestas_sim)}
    oficina.anotar(estado, momento, f"🏋️ Llegan de la pretemporada simulada ({progreso['dias']} días del "
                                    f"{progreso['desde']} al {progreso['hasta']}) con sus reglas y lecciones.")
    progreso["trasladada"] = True
    comun.guardar(f["progreso"], progreso)
    comun.guardar(config.ESTADO, estado)
    from agentes import web
    web.generar(estado, comun.leer(config.APUESTAS, []))
    return True


def main() -> None:
    p = argparse.ArgumentParser(description="Pretemporada simulada")
    p.add_argument("--dias", type=int, default=8, help="Días a simular en esta ejecución")
    p.add_argument("--desde", default=DESDE)
    p.add_argument("--hasta", default=HASTA)
    p.add_argument("--trasladar", action="store_true", help="Pasar lo aprendido a la oficina")
    p.add_argument("--activar", action="store_true", help="Desde ahora avanza sola en cada ejecución programada")
    p.add_argument("--solo-si-activa", action="store_true", help="Las ejecuciones programadas solo avanzan si está activada")
    args = p.parse_args()
    if args.trasladar:
        trasladar()
    else:
        avanzar(args.dias, args.desde, args.hasta, activar=args.activar, solo_si_activa=args.solo_si_activa)


if __name__ == "__main__":
    main()
